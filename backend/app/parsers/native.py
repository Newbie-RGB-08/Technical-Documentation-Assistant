"""NativeParser：不依赖 OCR/外部服务的基础格式解析（PRD 7.4.1）。

所有方法都按需 import 对应的第三方包，保证只安装基础依赖时也能加载本模块；
重度 OCR 依赖只在 RapidOCR 等独立解析器中导入。
"""
from __future__ import annotations

import csv
import io
import json
import re
import time
from pathlib import Path

from app.parsers.base import (
    DocumentParser,
    ParsedAsset,
    ParsedDocument,
    ParserError,
    UnsupportedFormatError,
)

SUPPORTED_EXTENSIONS = (
    ".md",
    ".markdown",
    ".txt",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".xls",
    ".csv",
    ".json",
)
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp")
MACRO_EXTENSIONS = {".docm", ".xlsm", ".pptm"}
REJECTED_MESSAGES = {
    ".doc": "首版不支持 .doc，请另存为 .docx",
    ".ppt": "首版不支持 .ppt，请另存为 .pptx",
}
_SUPPORTED_TEXT = ", ".join(SUPPORTED_EXTENSIONS)
_IMAGE_SUFFIXES = set(IMAGE_EXTENSIONS)


def validate_extension(filename: str) -> str:
    """校验扩展名并返回小写扩展名；不合法时抛出明确错误。"""
    ext = Path(str(filename or "")).suffix.lower()
    if not ext:
        raise UnsupportedFormatError(f"文件缺少扩展名；支持格式: {_SUPPORTED_TEXT}")
    if ext in REJECTED_MESSAGES:
        raise UnsupportedFormatError(REJECTED_MESSAGES[ext])
    if ext in MACRO_EXTENSIONS:
        raise UnsupportedFormatError("不支持宏格式，请另存为无宏格式")
    if ext not in SUPPORTED_EXTENSIONS and ext not in IMAGE_EXTENSIONS:
        raise UnsupportedFormatError(f"不支持的文件格式 {ext}；支持格式: {_SUPPORTED_TEXT}")
    return ext


def _pdf_reader(path: Path):
    """打开 PDF 并返回 reader；加密/0 页/坏页树统一转为 ParserError。"""
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ParserError("PDF 已加密，无法解析")
        pages = reader.pages
        if len(pages) <= 0:
            raise ParserError("PDF 没有页面")
        return reader
    except ParserError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParserError(f"PDF 打开失败: {type(exc).__name__}: {exc}") from exc


def pdf_page_count(path: Path | str) -> int:
    """返回 PDF 页数；预检失败时抛出 ParserError。"""
    return len(_pdf_reader(Path(path)).pages)


def pdf_text_chars_per_page(path: Path | str, sample: int = 20) -> float:
    """采样前 ``sample`` 页的文本层字符数/页。

    用于 auto 路由：数字 PDF 走 native，扫描 PDF 走 OCR。单页提取失败按 0 字符
    处理；加密、0 页、坏页树直接抛错。
    """
    reader = _pdf_reader(Path(path))
    pages = list(reader.pages)
    selected = pages[: max(1, int(sample))] if sample else pages
    total = 0
    for page in selected:
        try:
            total += len((page.extract_text() or "").strip())
        except Exception:  # noqa: BLE001 - 单页坏文本不阻断整份判定
            continue
    return total / max(1, len(selected))


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = data.decode(encoding)
            # 统一 IR 换行，避免 Windows 写入的 CRLF 进入切片/向量文本。
            return text.replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")


def _clean_cell(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return text.replace("|", "\\|").replace("\n", "<br>").strip()


def _markdown_table(rows: list[list[object]]) -> str:
    cleaned = [[_clean_cell(cell) for cell in row] for row in rows]
    cleaned = [row for row in cleaned if any(cell for cell in row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    if width == 0:
        return ""
    padded = [row + [""] * (width - len(row)) for row in cleaned]
    header = padded[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return "\n".join(lines)


def _heading_level(style_name: str | None) -> int | None:
    if not style_name:
        return None
    match = re.search(r"(?:heading|标题)\s*([1-6])", style_name, re.IGNORECASE)
    if match:
        return int(match.group(1))
    if style_name.strip().lower() in ("title", "标题"):
        return 1
    return None


def _safe_asset_dir(out_dir: Path) -> Path:
    assets = Path(out_dir).resolve() / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    return assets


class NativeParser(DocumentParser):
    name = "native"
    supported_extensions = SUPPORTED_EXTENSIONS

    def available(self) -> bool:
        return True

    def health(self) -> dict:
        started = time.perf_counter()
        return {
            "name": self.name,
            "available": True,
            "status": "ok",
            "detail": "built-in parser",
            "latency_ms": round((time.perf_counter() - started) * 1000, 1),
            "checked_at": time.time(),
            "formats": list(self.supported_extensions),
        }

    def parse(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        src = Path(src)
        ext = validate_extension(src.name)
        if ext in IMAGE_EXTENSIONS:
            raise UnsupportedFormatError("图片解析需要 RapidOCR/MinerU/PaddleX 引擎，native 不支持")
        out_dir = Path(out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        handlers = {
            ".md": self._parse_text,
            ".markdown": self._parse_text,
            ".txt": self._parse_text,
            ".html": self._parse_html,
            ".htm": self._parse_html,
            ".pdf": self._parse_pdf,
            ".docx": self._parse_docx,
            ".pptx": self._parse_pptx,
            ".xlsx": self._parse_xlsx,
            ".xls": self._parse_xls,
            ".csv": self._parse_csv,
            ".json": self._parse_json,
        }
        handler = handlers[ext]
        try:
            return handler(src, out_dir, params or {})
        except UnsupportedFormatError:
            raise
        except ImportError as exc:
            raise ParserError(
                f"缺少解析 {ext} 所需的依赖: {exc.name or exc}；请安装 backend/requirements.txt"
            ) from exc
        except ParserError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ParserError(f"{ext} 解析失败: {type(exc).__name__}: {exc}") from exc

    # ---- text / html / json ----
    def _parse_text(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        return ParsedDocument(markdown=_decode_bytes(src.read_bytes()), engine=self.name)

    def _parse_html(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        from bs4 import BeautifulSoup
        from markdownify import markdownify

        html = _decode_bytes(src.read_bytes())
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "noscript", "iframe"]):
            tag.decompose()
        markdown = markdownify(str(soup), heading_style="ATX")
        markdown = re.sub(r"\n{3,}", "\n\n", markdown).strip() + "\n"
        return ParsedDocument(markdown=markdown, engine=self.name)

    def _parse_json(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        value = json.loads(_decode_bytes(src.read_bytes()))
        body = json.dumps(value, ensure_ascii=False, indent=2)
        return ParsedDocument(markdown=f"```json\n{body}\n```\n", engine=self.name)

    # ---- pdf ----
    def _parse_pdf(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        from pypdf import PdfReader

        reader = PdfReader(str(src))
        pages: list[str] = []
        warnings: list[str] = []
        empty_pages = 0
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:  # noqa: BLE001
                text = ""
                warnings.append(f"第 {index} 页文本提取失败: {exc}")
            if not text.strip():
                empty_pages += 1
            pages.append(f"## 第 {index} 页\n\n{text.strip()}")
        if empty_pages and empty_pages == len(pages):
            warnings.append("PDF 未检测到文本层，可能需要 OCR")
        return ParsedDocument(
            markdown="\n\n".join(pages).strip() + ("\n" if pages else ""),
            engine=self.name,
            page_count=len(pages),
            warnings=warnings,
        )

    # ---- docx ----
    def _parse_docx(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(src))
        lines: list[str] = []
        assets: list[ParsedAsset] = []
        max_images = int(params.get("max_images", 300))
        image_index = 0
        warnings: list[str] = []

        def add_image(blob: bytes, suffix: str, caption: str = "") -> None:
            nonlocal image_index
            if image_index >= max_images:
                warnings.append(f"图片数量超过上限 {max_images}，已忽略剩余图片")
                return
            image_index += 1
            suffix = suffix.lower() if suffix.lower() in _IMAGE_SUFFIXES else ".png"
            name = f"image_{image_index:03d}{suffix}"
            assets_dir = _safe_asset_dir(out_dir)
            (assets_dir / name).write_bytes(blob)
            rel_path = f"assets/{name}"
            asset = ParsedAsset(
                kind="image",
                name=name,
                rel_path=rel_path,
                caption=caption,
            )
            assets.append(asset)
            lines.append(f"\n![{caption or name}]({rel_path})\n")

        body = document.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:tbl"):
                table = Table(child, document)
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                table_md = _markdown_table(rows)
                if table_md:
                    lines.append("\n" + table_md + "\n")
                continue
            if child.tag != qn("w:p"):
                continue
            paragraph = Paragraph(child, document)
            level = _heading_level(paragraph.style.name if paragraph.style else None)
            text = paragraph.text.strip()
            if text:
                lines.append(("#" * level + " " + text) if level else text)
                lines.append("")
            for blip in paragraph._p.xpath(".//a:blip"):
                rid = blip.get(qn("r:embed"))
                if not rid:
                    continue
                part = document.part.related_parts.get(rid)
                if part is None or not hasattr(part, "blob"):
                    continue
                suffix = Path(str(getattr(part, "partname", "image.png"))).suffix or ".png"
                add_image(part.blob, suffix)
        markdown = "\n".join(lines)
        markdown = re.sub(r"\n{3,}", "\n\n", markdown).strip() + "\n"
        return ParsedDocument(
            markdown=markdown,
            assets=assets,
            engine=self.name,
            warnings=warnings,
        )

    # ---- pptx ----
    def _parse_pptx(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(str(src))
        lines: list[str] = []
        assets: list[ParsedAsset] = []
        warnings: list[str] = []
        max_images = int(params.get("max_images", 300))
        image_index = 0

        def add_image(blob: bytes, suffix: str, page_no: int, caption: str = "") -> None:
            nonlocal image_index
            if image_index >= max_images:
                warnings.append(f"图片数量超过上限 {max_images}，已忽略剩余图片")
                return
            image_index += 1
            suffix = suffix.lower() if suffix.lower() in _IMAGE_SUFFIXES else ".png"
            name = f"image_{image_index:03d}{suffix}"
            assets_dir = _safe_asset_dir(out_dir)
            (assets_dir / name).write_bytes(blob)
            rel_path = f"assets/{name}"
            assets.append(
                ParsedAsset(
                    kind="image",
                    name=name,
                    rel_path=rel_path,
                    page_no=page_no,
                    caption=caption,
                )
            )
            lines.append(f"![{caption or name}]({rel_path})")

        def walk_shapes(shapes, page_no: int, *, skip_title: bool = False) -> None:
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk_shapes(shape.shapes, page_no)
                    continue
                if skip_title and shape == slide.shapes.title:
                    continue
                if getattr(shape, "has_table", False):
                    rows = [
                        [cell.text for cell in row.cells] for row in shape.table.rows
                    ]
                    table_md = _markdown_table(rows)
                    if table_md:
                        lines.append("")
                        lines.append(table_md)
                        lines.append("")
                    continue
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            lines.append(text)
                    continue
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                    try:
                        image = shape.image
                        add_image(image.blob, image.ext, page_no)
                    except Exception as exc:  # noqa: BLE001
                        warnings.append(f"第 {page_no} 页图片提取失败: {exc}")

        for page_no, slide in enumerate(presentation.slides, start=1):
            lines.append(f"## 第 {page_no} 页")
            title = slide.shapes.title
            if title is not None and getattr(title, "has_text_frame", False):
                title_text = title.text.strip()
                if title_text:
                    lines.append(f"### {title_text}")
            lines.append("")
            walk_shapes(slide.shapes, page_no, skip_title=True)
            lines.append("")
        markdown = "\n".join(lines)
        markdown = re.sub(r"\n{3,}", "\n\n", markdown).strip() + "\n"
        return ParsedDocument(
            markdown=markdown,
            assets=assets,
            engine=self.name,
            page_count=len(presentation.slides),
            warnings=warnings,
        )

    # ---- xlsx / xls / csv ----
    def _parse_xlsx(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        from openpyxl import load_workbook

        workbook = load_workbook(str(src), read_only=True, data_only=True)
        lines: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                lines.append(f"## {worksheet.title}")
                lines.append("")
                table = _markdown_table([list(row) for row in worksheet.iter_rows(values_only=True)])
                if table:
                    lines.append(table)
                lines.append("")
        finally:
            workbook.close()

        assets: list[ParsedAsset] = []
        warnings: list[str] = []
        # read_only 模式拿不到 _images；单独尽力打开一次，失败不影响表格解析。
        try:
            from openpyxl import load_workbook as load_images_workbook

            image_book = load_images_workbook(str(src), data_only=True, read_only=False)
            image_index = 0
            for worksheet in image_book.worksheets:
                for image in getattr(worksheet, "_images", []):
                    image_index += 1
                    blob = image._data() if callable(getattr(image, "_data", None)) else b""
                    if not blob:
                        continue
                    suffix = f".{str(getattr(image, 'format', 'png')).lower()}"
                    if suffix not in _IMAGE_SUFFIXES:
                        suffix = ".png"
                    name = f"image_{image_index:03d}{suffix}"
                    (_safe_asset_dir(out_dir) / name).write_bytes(blob)
                    rel_path = f"assets/{name}"
                    page_no = None
                    anchor = getattr(image, "anchor", None)
                    if anchor is not None and getattr(anchor, "_from", None) is not None:
                        page_no = anchor._from.row + 1
                    assets.append(
                        ParsedAsset(
                            kind="image",
                            name=name,
                            rel_path=rel_path,
                            page_no=page_no,
                            caption=worksheet.title,
                        )
                    )
                    lines.append(f"![{worksheet.title}]({rel_path})")
            image_book.close()
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"xlsx 图片提取失败（表格内容已保留）: {exc}")

        return ParsedDocument(
            markdown="\n".join(lines).strip() + "\n",
            assets=assets,
            engine=self.name,
            warnings=warnings,
        )

    def _parse_xls(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        import pandas as pd

        sheets = pd.read_excel(str(src), sheet_name=None, header=None)
        lines: list[str] = []
        for name, frame in sheets.items():
            lines.append(f"## {name}")
            lines.append("")
            table = _markdown_table(frame.where(frame.notna(), None).values.tolist())
            if table:
                lines.append(table)
            lines.append("")
        return ParsedDocument(markdown="\n".join(lines).strip() + "\n", engine=self.name)

    def _parse_csv(self, src: Path, out_dir: Path, params: dict) -> ParsedDocument:
        text = _decode_bytes(src.read_bytes())
        sample = text[:4096]
        delimiter = ","
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            pass
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        table = _markdown_table(rows)
        return ParsedDocument(markdown=(table + "\n") if table else "", engine=self.name)

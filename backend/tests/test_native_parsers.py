"""P2：NativeParser 各格式解析与拒绝策略。"""
from __future__ import annotations

import base64
import json
from io import BytesIO
from pathlib import Path

import pytest

from app.parsers.base import ParserError, UnsupportedFormatError
from app.parsers.native import NativeParser, validate_extension

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII="
)


def _parse(tmp_path: Path, filename: str, data: bytes | str):
    src = tmp_path / filename
    if isinstance(data, str):
        src.write_text(data, encoding="utf-8")
    else:
        src.write_bytes(data)
    out = tmp_path / "out"
    return NativeParser().parse(src, out, {})


def _make_pdf(path: Path, text: str = "Hello PDF text") -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET".encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as handle:
        writer.write(handle)


def test_text_markdown_and_front_matter(tmp_path):
    result = _parse(
        tmp_path,
        "a.md",
        "---\ntitle: Demo\ntags: [a]\n---\n\n# 标题\n\n正文",
    )
    assert "---\ntitle: Demo" in result.markdown
    assert "# 标题" in result.markdown
    assert result.engine == "native"


def test_text_gbk_fallback(tmp_path):
    result = _parse(tmp_path, "gbk.txt", "中文编码回退".encode("gbk"))
    assert "中文编码回退" in result.markdown


def test_html_removes_script_nav_and_converts_atx(tmp_path):
    html = """
    <html><body>
      <nav>导航不要</nav>
      <script>alert('x')</script>
      <h1>主标题</h1>
      <p>正文 <strong>加粗</strong></p>
      <footer>页脚不要</footer>
    </body></html>
    """
    result = _parse(tmp_path, "a.html", html)
    assert "# 主标题" in result.markdown
    assert "**加粗**" in result.markdown
    assert "导航不要" not in result.markdown
    assert "页脚不要" not in result.markdown
    assert "alert" not in result.markdown


def test_csv_and_json(tmp_path):
    result = _parse(tmp_path, "a.csv", "name,value\n甲,1\n乙,2\n")
    assert "| name | value |" in result.markdown
    assert "| 甲 | 1 |" in result.markdown

    result = _parse(tmp_path, "a.json", json.dumps({"标题": "值"}, ensure_ascii=False))
    assert "```json" in result.markdown
    assert '"标题": "值"' in result.markdown


def test_csv_gbk(tmp_path):
    result = _parse(tmp_path, "gbk.csv", "名称,数量\n苹果,3\n".encode("gbk"))
    assert "苹果" in result.markdown and "| 3 |" in result.markdown


def test_pdf_text_layer_and_empty_warning(tmp_path):
    pdf = tmp_path / "text.pdf"
    _make_pdf(pdf)
    result = NativeParser().parse(pdf, tmp_path / "out-pdf", {})
    assert "## 第 1 页" in result.markdown
    assert "Hello PDF text" in result.markdown
    assert result.page_count == 1
    assert not result.warnings

    from pypdf import PdfWriter

    empty = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with empty.open("wb") as handle:
        writer.write(handle)
    empty_result = NativeParser().parse(empty, tmp_path / "out-empty", {})
    assert any("OCR" in warning for warning in empty_result.warnings)


def test_docx_headings_table_and_image(tmp_path):
    from docx import Document
    from docx.shared import Inches

    png = tmp_path / "pixel.png"
    png.write_bytes(PNG_1X1)
    doc = Document()
    doc.add_heading("文档标题", level=1)
    doc.add_paragraph("正文段落")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列1"
    table.cell(0, 1).text = "列2"
    table.cell(1, 0).text = "值1"
    table.cell(1, 1).text = "值2"
    doc.add_picture(str(png), width=Inches(0.1))
    src = tmp_path / "a.docx"
    doc.save(src)
    out = tmp_path / "out-docx"
    result = NativeParser().parse(src, out, {})
    assert "# 文档标题" in result.markdown
    assert "正文段落" in result.markdown
    assert "| 列1 | 列2 |" in result.markdown
    assert "| 值1 | 值2 |" in result.markdown
    assert result.assets and result.assets[0].rel_path.startswith("assets/")
    assert (out / result.assets[0].rel_path).exists()
    assert f"]({result.assets[0].rel_path})" in result.markdown


def test_pptx_slide_text_table_and_image(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    png = tmp_path / "pixel.png"
    png.write_bytes(PNG_1X1)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "第一页标题"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text = "幻灯片正文"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    slide.shapes.add_picture(str(png), Inches(5), Inches(1), width=Inches(0.1))
    src = tmp_path / "a.pptx"
    presentation.save(src)
    result = NativeParser().parse(src, tmp_path / "out-pptx", {})
    assert "## 第 1 页" in result.markdown
    assert "### 第一页标题" in result.markdown
    assert "幻灯片正文" in result.markdown
    assert "| A | B |" in result.markdown
    assert result.assets and result.assets[0].page_no == 1


def test_xlsx_xls_and_csv_tables(tmp_path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["名称", "数量"])
    sheet.append(["苹果", 3])
    image_stream = BytesIO(PNG_1X1)
    sheet.add_image(XLImage(image_stream), "D2")
    xlsx = tmp_path / "a.xlsx"
    workbook.save(xlsx)
    result = NativeParser().parse(xlsx, tmp_path / "out-xlsx", {})
    assert "## 数据" in result.markdown
    assert "| 苹果 | 3 |" in result.markdown
    assert result.assets

    import xlwt

    book = xlwt.Workbook()
    old_sheet = book.add_sheet("旧表")
    old_sheet.write(0, 0, "A")
    old_sheet.write(0, 1, "B")
    old_sheet.write(1, 0, 1)
    old_sheet.write(1, 1, 2)
    xls = tmp_path / "a.xls"
    book.save(str(xls))
    old_result = NativeParser().parse(xls, tmp_path / "out-xls", {})
    assert "## 旧表" in old_result.markdown
    assert "| 1 | 2 |" in old_result.markdown


@pytest.mark.parametrize(
    ("filename", "fragment"),
    [
        ("a.doc", "另存为 .docx"),
        ("a.ppt", "另存为 .pptx"),
        ("a.docm", "宏格式"),
        ("a.xlsm", "宏格式"),
        ("a.pptm", "宏格式"),
        ("a.exe", "不支持"),
        ("a.zip", "不支持"),
    ],
)
def test_rejected_formats_are_rejected_before_output(tmp_path, filename, fragment):
    src = tmp_path / filename
    src.write_bytes(b"not a real document")
    out = tmp_path / f"out-{filename}"
    with pytest.raises(UnsupportedFormatError, match=fragment):
        NativeParser().parse(src, out, {})
    assert not out.exists()


def test_unknown_extension_and_image_native_rejection(tmp_path):
    with pytest.raises(UnsupportedFormatError):
        validate_extension("a.unknown")
    with pytest.raises(UnsupportedFormatError, match="RapidOCR"):
        NativeParser().parse(tmp_path / "a.png", tmp_path / "out-png", {})


def test_native_available_health_and_uppercase():
    parser = NativeParser()
    assert parser.available() is True
    health = parser.health()
    assert health["available"] is True and health["status"] == "ok"
    assert validate_extension("REPORT.PDF") == ".pdf"


def test_invalid_json_raises_parser_error(tmp_path):
    with pytest.raises(ParserError, match="json 解析失败"):
        _parse(tmp_path, "bad.json", "{not valid")
